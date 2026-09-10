"""Adapter for the paired arrival/departure daily-log layout.

Only normalized planning fields leave this module. Notes, aircraft identifiers,
auxiliary columns, worksheet names, and unrecognized source text are discarded.
"""

from dataclasses import replace
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Any
from uuid import UUID, uuid5
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.datetime import WINDOWS_EPOCH

from ramp_optimizer.config import OptimizerConfig, TeamWorkImportConfig
from ramp_optimizer.enums import IssueSeverity
from ramp_optimizer.models import Flight, OperationalDay
from ramp_optimizer.validation import validate_config, validate_operational_day
from ramp_optimizer_imports.enums import ImportType
from ramp_optimizer_imports.flight_models import (
    FlightCorrection,
    FlightReviewRow,
    FlightStatus,
    FlightTimePolicy,
)
from ramp_optimizer_imports.flight_normalization import (
    flight_number,
    location,
    operational_status,
    timestamp,
    validate_policy,
)
from ramp_optimizer_imports.models import ImportError, ImportPreview, ReviewIssue

HEADERS = (
    "FLT #",
    "CTY",
    "ETA",
    "NOTES",
    "OO",
    "A/C #",
    "CONF",
    "FLT #",
    "CTY",
    "ETD",
    "MST",
    "STATUS",
    "GATE",
)
FLIGHT_FIELDS = (
    "arrival_flight_number",
    "departure_flight_number",
    "arrival_time",
    "departure_time",
    "gate",
    "heavy",
)
CORRECTABLE = frozenset((*FLIGHT_FIELDS, "origin", "destination", "status", "excluded"))


def issue(
    code: str, row: int | None = None, field: str | None = None, *, warning: bool = False
) -> ReviewIssue:
    messages = {
        "INVALID_FLIGHT_VALUE": "Supply an explicit valid value for this flight field.",
        "CANCELLED_EXCLUDED": "Cancelled records are excluded from optimization demand.",
        "DUPLICATE_FLIGHT_ROW": "An identical active flight row must be explicitly excluded.",
        "FLIGHT_DOMAIN_VALIDATION": "The flight does not satisfy the optimizer flight rules.",
        "FLIGHT_DATE_OUTSIDE_WINDOW": "Use a date within the selected operational day or an adjacent calendar date.",
        "FORMULA_VALUE_UNAVAILABLE": "Replace the formula with an explicit reviewed planning value.",
        "EMPTY_FLIGHT_LOG": "At least one source flight record is required.",
        "MIDNIGHT_ROLLOVER": "The time-only departure rolls into the following calendar day.",
    }
    return ReviewIssue(
        code,
        IssueSeverity.WARNING if warning else IssueSeverity.ERROR,
        messages.get(code, "Review the indicated flight field."),
        row,
        field,
        not warning,
        "Correct the field or explicitly exclude the row." if row else None,
    )


def snapshot(preview: ImportPreview) -> OperationalDay:
    return OperationalDay(
        preview.operational_date,
        flights=tuple(
            row.flight
            for row in preview.flight_rows
            if not row.excluded and row.status != FlightStatus.CANCELLED
        ),
    )


def revalidate(preview: ImportPreview) -> ImportPreview:
    issues = [i for i in preview.source_issues if i.code != "WORKBOOK_DATE_MISMATCH"]
    active = []
    seen = set()
    if not preview.flight_rows:
        issues.append(issue("EMPTY_FLIGHT_LOG"))
    for row in preview.flight_rows:
        if row.status == FlightStatus.CANCELLED:
            issues.append(issue("CANCELLED_EXCLUDED", row.source_row, "status", warning=True))
        if row.excluded or row.status == FlightStatus.CANCELLED:
            continue
        active.append(row)
        for field in row.unresolved_fields:
            issues.append(issue("INVALID_FLIGHT_VALUE", row.source_row, field))
        for field in row.formula_fields:
            issues.append(issue("FORMULA_VALUE_UNAVAILABLE", row.source_row, field))
        for side in ("arrival", "departure"):
            if getattr(row, side + "_expected"):
                for suffix in ("_flight_number", "_time"):
                    if getattr(row.flight, side + suffix) is None:
                        issues.append(issue("INVALID_FLIGHT_VALUE", row.source_row, side + suffix))
        if row.status == FlightStatus.UNKNOWN or (
            row.status == FlightStatus.TERMINATING and row.departure_expected
        ):
            label = " / ".join(
                filter(None, (row.flight.arrival_flight_number, row.flight.departure_flight_number))
            )
            reason = {
                FlightStatus.UNKNOWN: "The workbook status is not recognized. Choose Normal if this flight should be staffed, Terminating for an arrival-only movement, or exclude it if no service is required.",
                FlightStatus.TERMINATING: "The workbook marks this flight as terminating, but also supplies a departure flight. Confirm whether it is a turn or an arrival-only movement before including it.",
            }[row.status]
            issues.append(
                ReviewIssue(
                    "STATUS_REQUIRES_REVIEW",
                    IssueSeverity.ERROR,
                    f"Flight {label}: {reason}",
                    row.source_row,
                    "status",
                    True,
                    "Use Correct to review the flight status and planned movement.",
                )
            )
        for field in ("arrival_time", "departure_time"):
            value = getattr(row.flight, field)
            if value is not None and abs((value.date() - preview.operational_date).days) > 1:
                issues.append(issue("FLIGHT_DATE_OUTSIDE_WINDOW", row.source_row, field))
        if row.flight in seen:
            issues.append(issue("DUPLICATE_FLIGHT_ROW", row.source_row, "excluded"))
        seen.add(row.flight)
    for item in validate_config(preview.config) + validate_operational_day(
        snapshot(preview), preview.config
    ):
        source_row = None
        if item.path.startswith("flights["):
            source_row = active[int(item.path.split("[", 1)[1].split("]", 1)[0])].source_row
        # Domain messages may include identifiers. Keep only the stable code/path.
        issues.append(
            ReviewIssue(
                item.code,
                IssueSeverity.ERROR,
                "The normalized input does not satisfy the optimizer rules.",
                source_row,
                item.path,
            )
        )
    return replace(preview, issues=tuple(issues))


class DailyFlightLogAdapter:
    import_type = ImportType.DAILY_FLIGHT_LOG
    schema_version = 1
    snapshot = staticmethod(snapshot)
    revalidate = staticmethod(revalidate)

    def parse(
        self,
        content: bytes,
        import_id: str,
        operational_date: date,
        policy: FlightTimePolicy,
        config: OptimizerConfig,
    ) -> ImportPreview:
        validate_policy(policy)
        try:
            workbook = load_workbook(
                BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
            try:
                candidates = []
                for sheet in workbook.worksheets:
                    sheet.reset_dimensions()
                    for index, cells in enumerate(sheet.iter_rows(max_row=50), 1):
                        labels = tuple(" ".join(str(c.value or "").upper().split()) for c in cells)
                        if labels[:13] == HEADERS and not any(labels[13:]):
                            candidates.append((sheet, index))
                if len(candidates) != 1:
                    raise ImportError("UNSUPPORTED_FLIGHT_LOG_LAYOUT")
                sheet, header = candidates[0]
                rows = []
                issues = []
                late = False
                for index, cells in enumerate(sheet.iter_rows(), 1):
                    values = [c.value for c in cells]
                    values += [None] * max(0, 13 - len(values))
                    if index < header:
                        continue
                    if index == header or not any(v not in (None, "") for v in values):
                        continue
                    if any(value not in (None, "") for value in values[13:]):
                        raise ImportError("UNSUPPORTED_FLIGHT_LOG_LAYOUT")
                    populated = [v for v in values if v not in (None, "")]
                    if len(populated) == 1 and isinstance(values[0], str):
                        label = values[0].strip().upper()
                        if label.startswith("LATE ARRIVALS"):
                            late = True
                            continue
                        if label.startswith(("NOTE:", "TOTAL", "SUBTOTAL")):
                            continue
                    row, rollover = _row(
                        values,
                        cells,
                        index,
                        import_id,
                        operational_date,
                        policy,
                        workbook.epoch,
                        late,
                    )
                    rows.append(row)
                    if rollover:
                        issues.append(
                            issue("MIDNIGHT_ROLLOVER", index, "departure_time", warning=True)
                        )
                preview = ImportPreview(
                    1,
                    operational_date,
                    (),
                    (),
                    (),
                    tuple(issues),
                    (),
                    config,
                    TeamWorkImportConfig(),
                    flight_rows=tuple(rows),
                    flight_policy=policy,
                )
                return revalidate(preview)
            finally:
                workbook.close()
        except ImportError:
            raise
        except (ValueError, KeyError, TypeError, IndexError, OSError, BadZipFile, ParseError):
            raise ImportError("INVALID_FLIGHT_LOG_WORKBOOK") from None
        except Exception:  # noqa: BLE001 -- parser boundary must not expose source values
            raise ImportError("FLIGHT_LOG_PARSE_FAILED", 500) from None

    def correct(
        self,
        preview: ImportPreview,
        corrections: tuple[FlightCorrection, ...],
        *,
        operational_date: date | None = None,
    ) -> ImportPreview:
        if preview.flight_policy is None:
            raise ImportError("INVALID_FLIGHT_TIME_POLICY")
        policy = preview.flight_policy
        if operational_date is not None:
            preview = replace(
                preview,
                operational_date=operational_date,
                operational_date_correction=(preview.operational_date, operational_date),
                source_issues=tuple(
                    i for i in preview.source_issues if i.code != "WORKBOOK_DATE_MISMATCH"
                ),
            )
        rows = {r.row_id: r for r in preview.flight_rows}
        seen = set()
        audit = []
        for correction in corrections:
            if correction.row_id not in rows or correction.row_id in seen:
                raise ImportError("INVALID_IMPORT_CORRECTIONS")
            seen.add(correction.row_id)
            row = rows[correction.row_id]
            changes = dict(correction.changes)
            if (
                not changes
                or len(changes) != len(correction.changes)
                or not changes.keys() <= CORRECTABLE
            ):
                raise ImportError("INVALID_IMPORT_CORRECTIONS")
            original = []
            flight_changes: dict[str, Any] = {}
            metadata: dict[str, Any] = {}
            for field, value in changes.items():
                before = getattr(row.flight if field in FLIGHT_FIELDS else row, field)
                original.append(
                    (field, before.isoformat() if isinstance(before, datetime) else before)
                )
                normalized: Any
                if field.endswith("_flight_number"):
                    normalized = flight_number(value)
                elif field.endswith("_time"):
                    normalized = timestamp(value, preview.operational_date, policy, WINDOWS_EPOCH)[
                        0
                    ]
                elif field in {"heavy", "excluded"}:
                    normalized = value if type(value) is bool else None
                elif field == "status":
                    try:
                        normalized = FlightStatus(value) if isinstance(value, str) else None
                    except (ValueError, TypeError):
                        normalized = None
                else:
                    normalized = location(value, airport=field in {"origin", "destination"})
                if normalized is None:
                    raise ImportError("INVALID_IMPORT_CORRECTIONS")
                (flight_changes if field in FLIGHT_FIELDS else metadata)[field] = normalized
            flight = replace(row.flight, **flight_changes)
            if "heavy" in changes:
                metadata["heavy_reviewed"] = True
            for side in ("arrival", "departure"):
                if side + "_flight_number" in changes or side + "_time" in changes:
                    metadata[side + "_expected"] = True
            rows[row.row_id] = replace(
                row,
                flight=flight,
                **metadata,
                unresolved_fields=tuple(f for f in row.unresolved_fields if f not in changes),
                formula_fields=tuple(f for f in row.formula_fields if f not in changes),
            )
            audit.append(FlightCorrection(row.row_id, correction.changes, tuple(original)))
        return revalidate(
            replace(
                preview,
                revision=preview.revision + 1,
                flight_rows=tuple(rows.values()),
                flight_corrections=tuple(audit),
            )
        )


def _row(values, cells, index, import_id, day, policy, epoch, late):
    mapping = {
        "arrival_flight_number": 0,
        "origin": 1,
        "arrival_time": 2,
        "departure_flight_number": 7,
        "destination": 8,
        "departure_time": 9,
        "status": 11,
        "gate": 12,
    }
    formula = tuple(
        field for field, col in mapping.items() if col < len(cells) and cells[col].data_type == "f"
    )
    clean = [
        None if col < len(cells) and cells[col].data_type == "f" else value
        for col, value in enumerate(values)
    ]
    arrival_expected = any(values[c] not in (None, "") for c in (0, 1, 2))
    departure_expected = any(values[c] not in (None, "") for c in (7, 8, 9))
    arrival_only = late or operational_status(clean[11]) == FlightStatus.TERMINATING
    onward = arrival_only and values[7] in (None, "") and values[9] not in (None, "")
    if arrival_only and values[7] in (None, ""):
        departure_expected = False
        formula = tuple(
            field for field in formula if field not in {"departure_time", "destination"}
        )
    arrival, arrival_dated = timestamp(clean[2], day, policy, epoch)
    departure, departure_dated = (
        timestamp(clean[9], day, policy, epoch) if departure_expected else (None, False)
    )
    rollover = bool(
        arrival and departure and departure < arrival and not arrival_dated and not departure_dated
    )
    if rollover:
        # Resolve the next wall-clock day again to detect DST gaps and folds.
        departure = timestamp(
            (departure + timedelta(days=1)).replace(tzinfo=None), day, policy, epoch
        )[0]
    # Fixed offsets preserve elapsed-time arithmetic across DST boundaries.
    arrival = datetime.fromisoformat(arrival.isoformat()) if arrival else None
    departure = datetime.fromisoformat(departure.isoformat()) if departure else None
    origin = location(clean[1], airport=True) if arrival_expected else None
    destination = location(clean[8], airport=True) if departure_expected else None
    gate = location(clean[12])
    unresolved = tuple(
        field
        for field, value, expected in (
            ("origin", origin, arrival_expected),
            ("destination", destination, departure_expected),
            ("gate", gate, bool(values[12])),
        )
        if expected and value is None
    )
    flight = Flight(
        flight_number(clean[0]) if arrival_expected else None,
        flight_number(clean[7]) if departure_expected else None,
        arrival,
        departure,
        gate,
        False,
    )
    return FlightReviewRow(
        str(uuid5(UUID(import_id), f"flight:{index}")),
        index,
        flight,
        arrival_expected,
        departure_expected,
        origin,
        destination,
        operational_status(clean[11]),
        notes_present=bool(values[3]),
        onward_time_present=onward,
        late_arrival_section=late,
        unresolved_fields=unresolved,
        formula_fields=formula,
        normalized_fields=tuple(
            field
            for field, col in mapping.items()
            if field.endswith("_flight_number")
            and flight_number(clean[col]) is not None
            and str(clean[col]) != flight_number(clean[col])
        ),
    ), rollover
