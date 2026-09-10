"""Entirely fictional workbooks; no external operational files are used."""

import json
from datetime import date, datetime, time
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import Workbook

from ramp_optimizer.config import OptimizerConfig
from ramp_optimizer.timing import derive_flight_operational_facts
from ramp_optimizer_imports.daily_flight_log import HEADERS, DailyFlightLogAdapter
from ramp_optimizer_imports.flight_models import FlightCorrection, FlightTimePolicy
from ramp_optimizer_imports.flight_normalization import flight_number, timestamp
from ramp_optimizer_imports.models import ImportError
from ramp_optimizer_imports.serialization import canonical_json, load_preview
from tests.test_import_workflow import client, service, storage  # noqa: F401

DAY = date(2032, 6, 15)
POLICY = FlightTimePolicy("America/Chicago", time(3), "ESTIMATED")


def workbook(rows=None, headers=HEADERS, extra_sheet=False):
    book = Workbook()
    sheet = book.active
    sheet.append(["FICTIONAL DAILY LOG 06/15/2032"])
    sheet.merge_cells("A1:M1")
    sheet.append(headers)
    for row in (
        rows
        if rows is not None
        else [[101, "ABC", "10:00", None, None, None, None, 102, "XYZ", "11:00", None, None, "A1"]]
    ):
        sheet.append(row)
    if extra_sheet:
        book.copy_worksheet(sheet)
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def parse(rows=None, **kwargs):
    return DailyFlightLogAdapter().parse(
        workbook(rows, **kwargs), str(uuid4()), DAY, POLICY, OptimizerConfig()
    )


def test_valid_turn_and_round_trip():
    preview = parse()
    assert preview.confirmation_eligible
    assert load_preview(canonical_json(preview)) == preview
    assert preview.flight_rows[0].flight.arrival_flight_number == "101"
    assert (
        derive_flight_operational_facts(
            preview.flight_rows[0].flight, preview.config
        ).flight_type.value
        == "TURN"
    )
    assert not preview.issues


def test_operational_date_correction_is_audited_without_shifting_times():
    preview = DailyFlightLogAdapter().parse(
        workbook(), str(uuid4()), date(2032, 6, 14), POLICY, OptimizerConfig()
    )
    assert preview.confirmation_eligible
    assert all(i.code != "WORKBOOK_DATE_MISMATCH" for i in preview.issues)
    revised = DailyFlightLogAdapter().correct(preview, (), operational_date=DAY)
    assert revised.operational_date == DAY
    assert revised.flight_rows[0].flight == preview.flight_rows[0].flight
    assert revised.operational_date_correction == (date(2032, 6, 14), DAY)
    assert load_preview(canonical_json(revised)) == revised


@pytest.mark.parametrize("second", ["101", "ZZ101"])
def test_conflicting_repeated_numbers_follow_domain_identity(second):
    preview = parse([["101", "ABC", "10:00"], [second, "XYZ", "12:00"]])
    assert not preview.confirmation_eligible


def test_explicit_reversed_dates_do_not_roll_forward():
    preview = parse(
        [
            [
                101,
                "ABC",
                datetime(2032, 6, 15, 23),
                None,
                None,
                None,
                None,
                102,
                "XYZ",
                datetime(2032, 6, 15, 1),
            ]
        ]
    )
    assert not preview.confirmation_eligible
    assert preview.flight_rows[0].flight.departure_time.date() == DAY


def test_opaque_auxiliary_fields_do_not_escape_preview(caplog):
    secret = "FICTIONAL-OPAQUE-REFERENCE-NEVER-RETAIN"
    preview = parse([[101, "ABC", "10:00", secret, secret, secret, secret]])
    assert secret not in canonical_json(preview)
    assert secret not in caplog.text
    assert preview.flight_rows[0].notes_present


def test_internal_parser_failure_is_sanitized(monkeypatch):
    def fail(*args, **kwargs):
        raise RuntimeError("FICTIONAL-SENSITIVE-PARSER-DETAIL")

    monkeypatch.setattr("ramp_optimizer_imports.daily_flight_log.load_workbook", fail)
    with pytest.raises(ImportError) as caught:
        parse()
    assert caught.value.status_code == 500
    assert str(caught.value) == "FLIGHT_LOG_PARSE_FAILED"


@pytest.mark.parametrize(
    "value,expected",
    [
        (" zz 00042 ", "ZZ42"),
        (42.0, "42"),
        ("42A", None),
        ("AA42/BB43", None),
        (True, None),
        (None, None),
    ],
)
def test_number_normalization(value, expected):
    assert flight_number(value) == expected


@pytest.mark.parametrize("kwargs", [{"headers": HEADERS[:-1]}, {"extra_sheet": True}])
def test_layout_rejected(kwargs):
    with pytest.raises(ImportError, match="UNSUPPORTED_FLIGHT_LOG_LAYOUT"):
        parse(**kwargs)


def test_nonadjacent_extra_column_is_not_silently_ignored():
    with pytest.raises(ImportError, match="UNSUPPORTED_FLIGHT_LOG_LAYOUT"):
        parse([[101, "ABC", "10:00", *([None] * 11), "FICTIONAL-EXTRA-COLUMN"]])


@pytest.mark.parametrize("number,expected", [(4999, False), (5000, False), (5001, True)])
def test_express_uses_domain(number, expected):
    config = OptimizerConfig(express_threshold=5000)
    preview = DailyFlightLogAdapter().parse(
        workbook([[number, "ABC", "10:00"]]), str(uuid4()), DAY, POLICY, config
    )
    assert (
        derive_flight_operational_facts(preview.flight_rows[0].flight, config).express is expected
    )


def test_late_arrivals_do_not_invent_departure():
    preview = parse(
        [
            ["LATE ARRIVALS / TERMINATING"],
            [101, "ABC", "23:00", None, None, None, None, None, "XYZ", "05:00", None, "TERM 102"],
        ]
    )
    row = preview.flight_rows[0]
    assert row.flight.departure_time is None
    assert row.onward_time_present
    assert preview.confirmation_eligible
    assert not preview.issues


@pytest.mark.parametrize("onward_value", [0, "=FICTIONAL_ONWARD_FORMULA()"])
def test_ignored_onward_time_is_not_required_planning_data(onward_value):
    preview = parse(
        [
            ["LATE ARRIVALS / TERMINATING"],
            [101, "ABC", "23:00", None, None, None, None, None, None, onward_value],
        ]
    )
    assert preview.confirmation_eligible
    assert preview.flight_rows[0].onward_time_present
    assert preview.flight_rows[0].flight.departure_time is None


def test_cancellation_and_unknown_status():
    preview = parse(
        [[101, "ABC", "10:00", None, None, None, None, None, None, None, None, "CANCELLED"]]
    )
    assert DailyFlightLogAdapter.snapshot(preview).flights == ()
    assert preview.confirmation_eligible
    preview = parse(
        [[101, "ABC", "10:00", None, None, None, None, None, None, None, None, "UNRECOGNIZED"]]
    )
    assert not preview.confirmation_eligible


def test_heavy_is_opt_in_without_a_review_warning():
    preview = parse()
    assert preview.flight_rows[0].flight.heavy is False
    assert not preview.issues
    revised = DailyFlightLogAdapter().correct(preview, (
        FlightCorrection(preview.flight_rows[0].row_id, (("heavy", True),)),
    ))
    assert revised.flight_rows[0].flight.heavy is True
    assert not revised.issues


def test_terminating_onward_time_never_extends_arrival_offload_work():
    preview = parse([
        [101, "ABC", "23:00", None, None, None, None, None, "XYZ", "05:00", None, "TERM 102"],
    ])
    assert not preview.issues
    flight = preview.flight_rows[0].flight
    assert flight.departure_time is None
    assert flight.departure_flight_number is None
    facts = derive_flight_operational_facts(flight, preview.config)
    assert facts.flight_type.value == "ARRIVAL_ONLY"
    assert (facts.work_end - flight.arrival_time).total_seconds() == 20 * 60


def test_aog_is_normal_and_preserves_planned_service():
    preview = parse([[None, None, None, None, None, None, None, 1134, "XYZ", "06:00", None, "AOG"]])
    assert preview.confirmation_eligible
    assert not preview.issues
    assert preview.flight_rows[0].status.value == "NORMAL"
    assert preview.flight_rows[0].flight.departure_flight_number == "1134"


def test_midnight_and_formula_correction():
    preview = parse([[101, "ABC", "23:00", None, None, None, None, 102, "XYZ", "01:00"]])
    flight = preview.flight_rows[0].flight
    assert flight.departure_time.date() == date(2032, 6, 16)
    preview = parse([[101, "ABC", "=PRIVATE_SOURCE()"]])
    assert not preview.confirmation_eligible
    assert "PRIVATE_SOURCE" not in canonical_json(preview)
    updated = DailyFlightLogAdapter().correct(
        preview,
        (
            FlightCorrection(
                preview.flight_rows[0].row_id,
                (("arrival_time", "2032-06-15T10:00:00-05:00"), ("heavy", True)),
            ),
        ),
    )
    assert updated.revision == 2
    assert updated.confirmation_eligible
    assert updated.flight_rows[0].flight.heavy
    assert updated.flight_corrections[0].original_values[0] == ("arrival_time", None)


def test_api_confirmation_idempotency_and_stale_revision(client):
    response = client.post(
        "/api/v1/imports/daily-flight-log",
        files={"workbook": ("fictional.xlsx", workbook())},
        data={
            "metadata": json.dumps(
                {
                    "operational_date": DAY.isoformat(),
                    "time_policy": {
                        "airport_timezone": POLICY.airport_timezone,
                        "operational_day_start": "03:00:00",
                        "planning_basis": "ESTIMATED",
                    },
                }
            )
        },
    )
    assert response.status_code == 201, response.text
    record = response.json()
    path = "/api/v1/imports/" + record["import_id"]
    response = client.post(
        path + "/corrections",
        json={
            "revision": 1,
            "flight_corrections": [
                {"row_id": record["preview"]["flight_rows"][0]["row_id"], "heavy": True}
            ],
        },
    )
    assert response.status_code == 200, response.text
    assert client.post(path + "/confirm", json={"revision": 1}).status_code == 409
    confirmed = client.post(path + "/confirm", json={"revision": 2})
    assert confirmed.status_code == 200, confirmed.text
    assert client.post(path + "/confirm", json={"revision": 2}).json() == confirmed.json()
    assert client.get(path + "/preview").json() == confirmed.json()
    assert (
        client.get(path + "/revisions/1").json()["preview"]["flight_rows"][0]["flight"]["heavy"]
        is False
    )
    assert client.get(path + "/revisions/99").status_code == 404
    day_id = confirmed.json()["confirmed_operational_day_id"]
    assert client.get("/api/v1/operational-days/" + day_id).json()["optimization_eligible"] is False
    assert client.post("/api/v1/operational-days/" + day_id + "/optimizations").status_code == 409


def test_sheet_and_aggregate_resource_limits():
    from ramp_optimizer_imports.safety import UploadLimits, validate_container

    with pytest.raises(ImportError, match="SUSPICIOUS_XLSX_ARCHIVE"):
        validate_container(workbook(extra_sheet=True), UploadLimits(max_worksheets=1))
    with pytest.raises(ImportError, match="SUSPICIOUS_XLSX_ARCHIVE"):
        validate_container(workbook(extra_sheet=True), UploadLimits(max_cells=30))


def test_migration_preserves_existing_employee_history(tmp_path, monkeypatch):
    from alembic import command
    from alembic.config import Config

    from ramp_optimizer import Employee
    from ramp_optimizer_imports.services import ImportService
    from ramp_optimizer_persistence.database import create_database_engine, make_session_factory
    from ramp_optimizer_persistence.imports import import_transactions
    from tests.import_fixtures import row
    from tests.import_fixtures import workbook as employee_workbook

    url = f"sqlite:///{(tmp_path / 'flight-migration.sqlite').as_posix()}"
    monkeypatch.setenv("RAMP_OPTIMIZER_DATABASE_URL", url)
    cfg = Config("alembic.ini")
    command.upgrade(cfg, "20260909_0002")
    engine = create_database_engine(url)
    imports = ImportService(import_transactions(make_session_factory(engine)))
    employee = imports.upload(
        employee_workbook([row(Date=DAY)]),
        "fictional.xlsx",
        "application/octet-stream",
        DAY,
        (Employee("SYN001", "Fictional Avery"),),
        OptimizerConfig(),
    )
    employee = imports.confirm(employee.import_id, 1)
    command.upgrade(cfg, "head")
    assert imports.get(employee.import_id) == employee
    command.downgrade(cfg, "20260909_0002")
    assert imports.get(employee.import_id) == employee
    command.upgrade(cfg, "head")
    flight = imports.upload_flights(
        workbook(), "fictional.xlsx", "application/octet-stream", DAY, POLICY, OptimizerConfig()
    )
    with pytest.raises(RuntimeError, match="Cannot downgrade"):
        command.downgrade(cfg, "20260909_0002")
    assert imports.get(flight.import_id) == flight
    engine.dispose()


@pytest.mark.parametrize(
    "value",
    [
        time(10),
        10 / 24,
        "10:00",
        "10:00 AM",
        datetime(2032, 6, 15, 10),
        "2032-06-15T10:00:00-05:00",
    ],
)
def test_time_representations(value):
    from openpyxl.utils.datetime import WINDOWS_EPOCH

    result, _ = timestamp(value, DAY, POLICY, WINDOWS_EPOCH)
    assert result.hour == 10
    assert result.date() == DAY
    assert result.utcoffset().total_seconds() == -18000


@pytest.mark.parametrize("value", ["2032-03-14T02:30:00", "2032-11-07T01:30:00"])
def test_dst_ambiguous_wall_time_requires_explicit_offset(value):
    from openpyxl.utils.datetime import WINDOWS_EPOCH

    assert timestamp(value, DAY, POLICY, WINDOWS_EPOCH)[0] is None


def test_duplicates_and_exclusion_revalidation():
    preview = parse([[101, "ABC", "10:00"], [101, "ABC", "10:00"]])
    assert not preview.confirmation_eligible
    revised = DailyFlightLogAdapter().correct(
        preview, (FlightCorrection(preview.flight_rows[1].row_id, (("excluded", True),)),)
    )
    assert revised.confirmation_eligible
    assert len(DailyFlightLogAdapter.snapshot(revised).flights) == 1


def test_departure_only_blank_notes_and_bad_required_number():
    preview = parse(
        [
            [],
            ["NOTE: fictional decorative text"],
            [None, None, None, None, None, None, None, 102, "XYZ", "11:00"],
        ]
    )
    assert len(preview.flight_rows) == 1
    assert preview.confirmation_eligible
    assert preview.flight_rows[0].flight.arrival_time is None
    assert not parse([[None, "ABC", "10:00"]]).confirmation_eligible


def test_confirm_rollback(service, storage, monkeypatch):
    from sqlalchemy import func, select

    from ramp_optimizer_persistence.imports import SQLImportRepository
    from ramp_optimizer_persistence.models import OperationalDayRow

    record = service.upload_flights(
        workbook(), "fictional.xlsx", "application/octet-stream", DAY, POLICY, OptimizerConfig()
    )

    def fail(*args):
        raise RuntimeError("synthetic failure")

    monkeypatch.setattr(SQLImportRepository, "confirm", fail)
    with pytest.raises(RuntimeError, match="synthetic failure"):
        service.confirm(record.import_id, 1)
    with storage[1]() as session:
        assert session.scalar(select(func.count()).select_from(OperationalDayRow)) == 0
    assert service.get(record.import_id).confirmed_at is None


def test_combined_inputs_use_domain_only(service, storage, monkeypatch):
    from ramp_optimizer import Employee, optimize_flight_assignments
    from ramp_optimizer_api.services import PersistenceService
    from tests.import_fixtures import row
    from tests.import_fixtures import workbook as employee_workbook

    employee = service.upload(
        employee_workbook([row(Date=DAY)]),
        "fictional.xlsx",
        "application/octet-stream",
        DAY,
        (Employee("SYN001", "Fictional Avery"),),
        OptimizerConfig(),
    )
    employee = service.confirm(employee.import_id, 1)
    flight = service.upload_flights(
        workbook(), "fictional.xlsx", "application/octet-stream", DAY, POLICY, OptimizerConfig()
    )
    flight = service.confirm(flight.import_id, 1)
    assert not service.readiness(employee.confirmed_operational_day_id)["optimization_eligible"]
    assert not service.readiness(flight.confirmed_operational_day_id)["optimization_eligible"]
    combined = service.combine(employee.import_id, flight.import_id)
    assert combined["optimization_eligible"]
    assert service.combine(employee.import_id, flight.import_id) == combined

    def forbidden(*args, **kwargs):
        raise AssertionError("Optimization must not read workbooks")

    monkeypatch.setattr("ramp_optimizer_imports.daily_flight_log.load_workbook", forbidden)
    monkeypatch.setattr(
        "ramp_optimizer_imports.employee_schedule.import_teamwork_schedule", forbidden
    )
    calls = []

    def optimizer(day, config):
        calls.append(day)
        assert len(day.flights) == 1 and len(day.employee_shifts) == 1
        assert day.employee_shifts[0].start.utcoffset() == day.flights[0].arrival_time.utcoffset()
        return optimize_flight_assignments(day, config)

    persistence = PersistenceService(storage[1], api_version="1", optimizer=optimizer)
    result = persistence.optimize_operational_day(combined["operational_day_id"])
    assert result.operational_day_id == combined["operational_day_id"]
    assert len(calls) == 1


def test_time_basis_and_auxiliary_time_never_override_eta():
    policy = FlightTimePolicy("America/Chicago", time(3), "SCHEDULED")
    preview = DailyFlightLogAdapter().parse(
        workbook([[101, "ABC", "10:00", None, None, None, None, None, None, None, "15:00"]]),
        str(uuid4()),
        DAY,
        policy,
        OptimizerConfig(),
    )
    assert preview.flight_policy.planning_basis == "SCHEDULED"
    assert preview.flight_rows[0].flight.arrival_time.hour == 10


def test_deterministic_preview_order():
    adapter = DailyFlightLogAdapter()
    import_id = str(uuid4())
    content = workbook([[103, "XYZ", "12:00"], [101, "ABC", "10:00"]])
    first = adapter.parse(content, import_id, DAY, POLICY, OptimizerConfig())
    assert adapter.parse(content, import_id, DAY, POLICY, OptimizerConfig()) == first
    assert [row.source_row for row in first.flight_rows] == [3, 4]


def test_arrival_aog_is_also_normal():
    preview = parse([[101, "ABC", "10:00", None, None, None, None, None, None, None, None, "AOG"]])
    assert preview.confirmation_eligible
    corrected = DailyFlightLogAdapter().correct(
        preview, (FlightCorrection(preview.flight_rows[0].row_id, (("status", "NORMAL"),)),)
    )
    assert corrected.confirmation_eligible


@pytest.mark.parametrize(
    "filename,content,status",
    [("fictional.xlsm", b"anything", 415), ("fictional.xlsx", b"not a workbook", 422)],
)
def test_flight_upload_rejects_bad_files(client, filename, content, status):
    response = client.post(
        "/api/v1/imports/daily-flight-log",
        files={"workbook": (filename, content)},
        data={
            "metadata": json.dumps(
                {
                    "operational_date": DAY.isoformat(),
                    "time_policy": {
                        "airport_timezone": "America/Chicago",
                        "operational_day_start": "03:00",
                        "planning_basis": "ESTIMATED",
                    },
                }
            )
        },
    )
    assert response.status_code == status


def test_combination_requires_confirmed_types(client):
    response = client.post(
        "/api/v1/imports/combine",
        json={"employee_import_id": str(uuid4()), "flight_import_id": str(uuid4())},
    )
    assert response.status_code == 404
